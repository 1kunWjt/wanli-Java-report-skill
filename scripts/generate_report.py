#!/usr/bin/env python3
"""Generate a Java lab report .docx from JSON data.

Usage:
    python generate_report.py data.json [output.docx]
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SCRIPT_DIR = Path(__file__).resolve().parent
CONSOLE_SCRIPT = SCRIPT_DIR / "generate_console.py"


def set_run_font(run, east_asia: str, latin: str | None = None, size: float | None = None, bold: bool | None = None):
    latin = latin or east_asia
    run.font.name = latin
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_para(
    doc_or_cell,
    text: str = "",
    *,
    east_asia: str = "宋体",
    latin: str = "Times New Roman",
    size: float = 12,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    space_before: float = 0,
    space_after: float = 6,
    first_line_indent: float | None = None,
    line_spacing: float = 1.25,
):
    p = doc_or_cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    if text:
        run = p.add_run(text)
        set_run_font(run, east_asia=east_asia, latin=latin, size=size, bold=bold)
    return p


def add_heading_cn(doc_or_cell, text: str):
    return add_para(
        doc_or_cell,
        text,
        east_asia="黑体",
        latin="Times New Roman",
        size=14,
        bold=True,
        space_before=12,
        space_after=6,
    )


def add_body(doc_or_cell, text: str, indent: bool = False):
    return add_para(
        doc_or_cell,
        text,
        east_asia="宋体",
        latin="Times New Roman",
        size=12,
        bold=False,
        first_line_indent=0.74 if indent else None,
        space_after=4,
        line_spacing=1.3,
    )


def add_code_block(doc_or_cell, code: str):
    for line in (code or "").splitlines() or [""]:
        p = add_para(
            doc_or_cell,
            line if line else " ",
            east_asia="宋体",
            latin="Consolas",
            size=10.5,
            bold=False,
            space_after=0,
            line_spacing=1.15,
        )
        p.paragraph_format.left_indent = Cm(0.5)


def set_table_borders(table, size: int = 12, color: str = "000000"):
    """Apply borders to the table and every cell (more reliable across Word/WPS)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    tblPr.append(borders)

    edge = {"val": "single", "sz": size, "color": color}
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=edge, left=edge, bottom=edge, right=edge)


def set_table_no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    tblPr.append(borders)


def set_cell_border(cell, **kwargs):
    """Set cell border. kwargs: top/left/bottom/right = {val, sz, color}."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        if edge not in kwargs:
            continue
        edge_data = kwargs.get(edge)
        tag = f"w:{edge}"
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        element.set(qn("w:val"), edge_data.get("val", "single"))
        element.set(qn("w:sz"), str(edge_data.get("sz", 12)))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), edge_data.get("color", "000000"))


def clear_cell(cell):
    # keep one empty paragraph
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    cell.add_paragraph()


def add_field_line(
    container,
    label: str,
    value: str,
    *,
    label_size: float = 14,
    value_size: float = 14,
    pad: str = "　　",
    underline_value: bool = True,
    trailing_underline: str = "________________",
    space_after: float = 8,
):
    """课程名称：______ value ______ style line with underlined value."""
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.3

    label_run = p.add_run(label)
    set_run_font(label_run, east_asia="宋体", latin="Times New Roman", size=label_size, bold=True)

    # leading blank underline (like form paper)
    lead = p.add_run("__________")
    set_run_font(lead, east_asia="宋体", latin="Times New Roman", size=value_size, bold=True)
    lead.font.underline = True

    pad_run = p.add_run(pad)
    set_run_font(pad_run, east_asia="宋体", latin="Times New Roman", size=value_size, bold=True)

    val_run = p.add_run(value or "")
    set_run_font(val_run, east_asia="宋体", latin="Times New Roman", size=value_size, bold=True)
    if underline_value:
        val_run.font.underline = True

    trail = p.add_run(trailing_underline)
    set_run_font(trail, east_asia="宋体", latin="Times New Roman", size=value_size, bold=True)
    trail.font.underline = True
    return p


def add_score_box(cell):
    """成绩/教师 bordered box content (tall enough to match two left rows)."""
    clear_cell(cell)
    border = {"val": "single", "sz": 14, "color": "000000"}
    set_cell_border(cell, top=border, left=border, bottom=border, right=border)

    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(18)
    p1.paragraph_format.space_after = Pt(18)
    r1 = p1.add_run("成绩：________")
    set_run_font(r1, east_asia="宋体", latin="Times New Roman", size=14, bold=True)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(18)
    p2.paragraph_format.space_after = Pt(18)
    r2 = p2.add_run("教师：________")
    set_run_font(r2, east_asia="宋体", latin="Times New Roman", size=14, bold=True)

    # extra blank to visually match left two-row height
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_before = Pt(2)
    p3.paragraph_format.space_after = Pt(2)
    r3 = p3.add_run("　")
    set_run_font(r3, east_asia="宋体", latin="Times New Roman", size=14, bold=True)


def add_screenshots(container, paths: list[str], placeholder: str = "（请在此插入运行结果截图）"):
    inserted = 0
    for raw in paths or []:
        path = Path(raw)
        if path.is_file():
            try:
                p = container.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(path), width=Cm(12))
                inserted += 1
            except Exception as exc:  # noqa: BLE001
                add_body(container, f"（截图插入失败：{path.name}：{exc}）")
        else:
            add_body(container, f"（未找到截图文件：{raw}）")
    if inserted == 0:
        add_body(container, placeholder)


def normalize_multiline(text: str) -> str:
    return (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def infer_class_and_args(code: str, run_hint: str | None) -> tuple[str, str]:
    """Return (class_name, display_command)."""
    code = code or ""
    m = re.search(r"\bclass\s+([A-Za-z_]\w*)", code)
    class_name = m.group(1) if m else "Main"
    if run_hint:
        hint = run_hint.strip()
        if hint.startswith("java "):
            return class_name, hint
        return class_name, f"java {class_name} {hint}".strip()
    return class_name, f"java {class_name}"


def guess_expected_output(code: str) -> str:
    lines = []
    for m in re.finditer(r"System\.out\.print(?:ln)?\((.*?)\);", code or "", re.S):
        arg = m.group(1)
        lits = re.findall(r'"([^"]*)"', arg)
        if lits:
            lines.append("".join(lits) + ("…" if "+" in arg else ""))
    return "\n".join(lines)


def ensure_console_screenshots(item: dict, index: int, dest_dir: Path) -> list[str]:
    """Return user screenshots, or auto-generate a CMD-style PNG for this problem."""
    paths = [str(p) for p in (item.get("screenshots") or []) if p]
    if paths:
        return paths

    code = normalize_multiline(item.get("code"))
    if not code:
        return []
    class_name, command = infer_class_and_args(code, normalize_multiline(item.get("run_hint")))
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"console-{index}.png"

    with tempfile.TemporaryDirectory(prefix="lab-code-") as tmp:
        src = Path(tmp) / f"{class_name}.java"
        src.write_text(code, encoding="utf-8")
        cmd = [
            sys.executable,
            str(CONSOLE_SCRIPT),
            str(dest),
            "--cmd",
            command,
            "--code",
            str(src),
        ]
        try:
            subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=40,
                check=False,
            )
        except Exception:
            pass

    if dest.is_file():
        return [str(dest)]

    expected = (
        normalize_multiline(item.get("expected_output"))
        or guess_expected_output(code)
        or "(无输出)"
    )
    try:
        sys.path.insert(0, str(SCRIPT_DIR))
        from generate_console import generate_console_image

        generate_console_image(dest, command, expected)
        return [str(dest)] if dest.is_file() else []
    except Exception:
        return []


def _add_underline_run(p, text: str, size: float = 14):
    run = p.add_run(text)
    set_run_font(run, east_asia="宋体", latin="Times New Roman", size=size, bold=True)
    run.font.underline = True
    return run


def _add_plain_run(p, text: str, size: float = 14):
    run = p.add_run(text)
    set_run_font(run, east_asia="宋体", latin="Times New Roman", size=size, bold=True)
    return run


def add_cover(doc: Document, data: dict):
    """Title + score box + underlined info fields (matches template header)."""
    school = data.get("school_name") or "浙江万里学院"
    course = data.get("course_name") or ""
    lab_name = data.get("lab_name") or ""
    class_name = data.get("class_name") or ""
    student_name = data.get("student_name") or ""
    student_id = data.get("student_id") or ""
    lab_date = data.get("lab_date") or ""

    # 大标题
    add_para(
        doc,
        f"{school}实验报告",
        east_asia="黑体",
        latin="Times New Roman",
        size=26,
        bold=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=18,
        space_after=28,
    )

    # 左侧字段 + 右侧成绩/教师方框（固定列宽，单行填空场禁止折行）
    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    set_table_no_borders(header)
    left, right = header.rows[0].cells
    left.width = Cm(11.5)
    right.width = Cm(4.5)

    tblGrid = header._tbl.tblGrid
    for grid_col, width in zip(tblGrid.findall(qn("w:gridCol")), (Cm(11.5), Cm(4.5))):
        grid_col.set(qn("w:w"), str(int(width.twips)))

    clear_cell(left)

    def fill_field(label: str, value: str, total_blank: int = 18) -> str:
        """Build one underlined blank field that fits a single line (no wrap)."""
        # fullwidth slots left for value+padding after label; keep total modest
        # label already takes space; blank field uses underscores (narrower than 全角空格)
        value = value or ""
        # Use ASCII underscores — much narrower, less likely to wrap
        lead_n = 4
        # remaining underscores after value; always >= 3 so it looks like a form
        trail_n = max(3, total_blank - lead_n - len(value))
        return ("_" * lead_n) + value + ("_" * trail_n)

    # 课程名称：________ 值 ________
    p = left.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.35
    # disable wrap by keeping the underlined field short enough for 11.5cm
    _add_plain_run(p, "课程名称：")
    _add_underline_run(p, fill_field("", course, total_blank=16))

    p2 = left.add_paragraph()
    p2.paragraph_format.space_before = Pt(8)
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.line_spacing = 1.35
    _add_plain_run(p2, "实验名称：")
    _add_underline_run(p2, (lab_name or "") + "____")

    # 右侧成绩/教师方框：与左侧两行等高
    add_score_box(right)
    tcPr = right._tc.get_or_add_tcPr()
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    tcPr.append(v_align)

    # 班级行：四个字段连排，值加下划线
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p3.paragraph_format.space_before = Pt(16)
    p3.paragraph_format.space_after = Pt(18)
    p3.paragraph_format.line_spacing = 1.4
    # 13pt keeps the packed class/name/id/date row on one line under A4
    meta_size = 13

    def _seg(text: str, underline: bool = False):
        run = p3.add_run(text)
        set_run_font(run, east_asia="宋体", latin="Times New Roman", size=meta_size, bold=True)
        if underline:
            run.font.underline = True

    _seg("专业班级：")
    _seg(class_name, underline=True)
    _seg("姓名：")
    _seg(student_name, underline=True)
    _seg("学号：")
    _seg(student_id, underline=True)
    _seg("实验日期：")
    _seg(lab_date, underline=True)


def add_body_box(doc: Document, data: dict, out_assets_dir: Path):
    """Big bordered box containing 一～五 sections (matches template)."""
    purpose = normalize_multiline(data.get("purpose"))
    content_summary = normalize_multiline(data.get("content_summary"))
    summary = normalize_multiline(data.get("summary"))
    problems = data.get("problems") or []

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, size=12, color="000000")
    cell = table.cell(0, 0)
    clear_cell(cell)

    # 一、实验目的
    add_heading_cn(cell, "一、实验目的：")
    for line in purpose.splitlines() or [purpose]:
        add_body(cell, line, indent=True)

    # 二、实验内容
    add_heading_cn(cell, "二、实验内容：")
    if content_summary:
        for line in content_summary.splitlines():
            add_body(cell, line)
    else:
        for idx, item in enumerate(problems, 1):
            add_body(cell, f"实验{idx}：{normalize_multiline(item.get('title'))}")

    # 三、实验步骤
    add_heading_cn(cell, "三、实验步骤：")
    for idx, item in enumerate(problems, 1):
        thinking = normalize_multiline(item.get("thinking"))
        title = normalize_multiline(item.get("title"))
        label = f"实验{idx}："
        lines = thinking.splitlines() if thinking else []
        if not lines:
            add_body(cell, f"{label}{title}")
        elif lines[0].startswith("实验") and "：" in lines[0]:
            for line in lines:
                add_body(cell, line)
        else:
            add_body(cell, f"{label}{lines[0]}")
            for line in lines[1:]:
                add_body(cell, line)

    # 四、实验结果
    add_heading_cn(cell, "四、实验结果：")
    for idx, item in enumerate(problems, 1):
        title = normalize_multiline(item.get("title")) or f"实验{idx}"
        if not title.startswith("实验"):
            title = f"实验{idx}：{title}"
        add_body(cell, title)
        add_body(cell, "1）源代码")
        add_code_block(cell, normalize_multiline(item.get("code")))
        add_body(cell, "2）运行结果（截图）")
        shots = ensure_console_screenshots(item, idx, out_assets_dir)
        add_screenshots(cell, shots)
        add_para(cell, "", space_after=6)

    # 五、实验总结
    add_heading_cn(cell, "五、实验总结：")
    if summary:
        for line in summary.splitlines():
            add_body(cell, line, indent=True)
    else:
        add_body(cell, "（请补充实验总结）", indent=True)


def build_document(data: dict, out_assets_dir: Path | None = None) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)

    add_cover(doc, data)
    add_body_box(doc, data, out_assets_dir or Path(tempfile.mkdtemp(prefix="lab-shots-")))
    return doc


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("Usage: python generate_report.py data.json [output.docx]", file=sys.stderr)
        return 2

    data_path = Path(argv[1])
    if not data_path.is_file():
        print(f"JSON not found: {data_path}", file=sys.stderr)
        return 1

    with data_path.open("r", encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, dict):
        print("JSON root must be an object", file=sys.stderr)
        return 1

    out_path = Path(argv[2]) if len(argv) > 2 else Path("实验报告.docx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_assets_dir = out_path.parent / f"{out_path.stem}-截图"

    doc = build_document(data, out_assets_dir=out_assets_dir)
    doc.save(str(out_path))

    reopened = Document(str(out_path))
    text_len = sum(len(p.text) for p in reopened.paragraphs)
    for t in reopened.tables:
        for row in t.rows:
            for cell in row.cells:
                text_len += sum(len(p.text) for p in cell.paragraphs)
    if text_len < 50:
        print(f"WARN: generated document looks too short ({text_len} chars)", file=sys.stderr)

    print(f"OK: {out_path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
